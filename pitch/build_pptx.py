# -*- coding: utf-8 -*-
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
ROOT = Path(__file__).resolve().parent
SHOTS = ROOT / "screenshots"
OUT = ROOT / "BOAR-prezentaciya.pptx"

# Brand
PINE = RGBColor(0x1C, 0x3D, 0x38)
PINE_DARK = RGBColor(0x0F, 0x28, 0x24)
COPPER = RGBColor(0x9A, 0x6B, 0x3D)
PAPER = RGBColor(0xF3, 0xEE, 0xE4)
SHEET = RGBColor(0xFA, 0xF7, 0xF1)
INK = RGBColor(0x17, 0x24, 0x1F)
MUTED = RGBColor(0x66, 0x75, 0x6E)
WHITE = RGBColor(0xFF, 0xFC, 0xF7)
BAD = RGBColor(0x8A, 0x3A, 0x2A)
OK = RGBColor(0x2B, 0x5C, 0x45)


def set_run(run, size=18, bold=False, color=INK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_bg(slide, color=PAPER):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    fill_shape(shape, color)
    # send to back
    spTree = slide.shapes._spTree
    sp = shape._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def add_footer(slide, page, total=12):
    add_textbox(slide, Inches(0.5), Inches(7.05), Inches(8), Inches(0.3), "БОАР · новый сайт общества", size=11, color=MUTED)
    add_textbox(slide, Inches(11.2), Inches(7.05), Inches(1.6), Inches(0.3), f"{page} / {total}", size=11, color=MUTED, align=PP_ALIGN.RIGHT)


def fit_image(slide, path, left, top, max_w, max_h):
    if not path.exists():
        return None
    pic = slide.shapes.add_picture(str(path), left, top)
    # scale to fit
    w, h = pic.width, pic.height
    scale = min(max_w / w, max_h / h)
    pic.width = int(w * scale)
    pic.height = int(h * scale)
    return pic


def main():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    total = 12

    # 1 Title
    s = prs.slides.add_slide(blank)
    add_bg(s, PINE_DARK)
    add_textbox(s, Inches(0.7), Inches(1.6), Inches(11), Inches(0.4), "ПРЕЗЕНТАЦИЯ ДЛЯ РУКОВОДСТВА ОБЩЕСТВА", size=14, bold=True, color=COPPER)
    add_textbox(s, Inches(0.7), Inches(2.2), Inches(11.5), Inches(2), "Новый сайт БОАР\nкак рабочий инструмент врача", size=40, bold=True, color=WHITE, font="Georgia")
    add_textbox(s, Inches(0.7), Inches(4.6), Inches(10), Inches(1), "Заседания, документы, обучение, вступление и кабинет члена —\nв одном современном продукте.", size=18, color=RGBColor(0xC5, 0xD4, 0xCE))
    add_textbox(s, Inches(0.7), Inches(6.5), Inches(10), Inches(0.4), "Сентябрь 2026 · макет для согласования", size=12, color=MUTED)

    # 2 Why
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.3), "ЗАЧЕМ МЕНЯТЬ", size=12, bold=True, color=COPPER)
    add_textbox(s, Inches(0.6), Inches(0.7), Inches(12), Inches(0.8), "Что нужно профессору и практикующему врачу", size=28, bold=True, color=INK, font="Georgia")
    items = [
        ("01", "Быстро найти заседание", "Дата, зал, программа — без поиска по ленте новостей."),
        ("02", "Открыть нужный документ", "Протокол, санэпид, ФАР — по разделам."),
        ("03", "Показать общество коллегам", "Сайт, который не стыдно открыть на лекции и съезде."),
        ("04", "Принять нового члена", "Понятный путь заявки вместо письма «куда-нибудь»."),
    ]
    for i, (num, title, text) in enumerate(items):
        x = Inches(0.6 + (i % 2) * 6.2)
        y = Inches(1.8 + (i // 2) * 2.3)
        card = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(5.9), Inches(2.0))
        fill_shape(card, SHEET)
        card.line.color.rgb = RGBColor(0xD8, 0xD0, 0xC4)
        add_textbox(s, x + Inches(0.25), y + Inches(0.25), Inches(5.3), Inches(0.3), num, size=14, bold=True, color=COPPER)
        add_textbox(s, x + Inches(0.25), y + Inches(0.6), Inches(5.3), Inches(0.4), title, size=18, bold=True, color=INK, font="Georgia")
        add_textbox(s, x + Inches(0.25), y + Inches(1.1), Inches(5.3), Inches(0.7), text, size=14, color=MUTED)
    add_footer(s, 2, total)

    # 3 Home screenshot
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_textbox(s, Inches(0.5), Inches(0.25), Inches(12), Inches(0.3), "ГЛАВНАЯ", size=12, bold=True, color=COPPER)
    add_textbox(s, Inches(0.5), Inches(0.55), Inches(12), Inches(0.5), "Ближайшее заседание — сразу в первом экране", size=24, bold=True, color=INK, font="Georgia")
    fit_image(s, SHOTS / "01-home.png", Inches(0.8), Inches(1.2), Inches(11.7), Inches(5.5))
    add_footer(s, 3, total)

    # 4 Old vs new
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.3), "СРАВНЕНИЕ", size=12, bold=True, color=COPPER)
    add_textbox(s, Inches(0.6), Inches(0.7), Inches(12), Inches(0.6), "Старый сайт → новый продукт", size=28, bold=True, color=INK, font="Georgia")
    left = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.6), Inches(5.8), Inches(5.0))
    fill_shape(left, SHEET)
    left.line.color.rgb = BAD
    right = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.9), Inches(1.6), Inches(5.8), Inches(5.0))
    fill_shape(right, SHEET)
    right.line.color.rgb = OK
    add_textbox(s, Inches(0.85), Inches(1.8), Inches(5.2), Inches(0.4), "Было (bsaer.org)", size=18, bold=True, color=BAD, font="Georgia")
    add_textbox(s, Inches(7.15), Inches(1.8), Inches(5.2), Inches(0.4), "Стало (новый макет)", size=18, bold=True, color=OK, font="Georgia")
    old = "• Новостная лента как главный экран\n• Документы «где-то в разделах»\n• Съезды и заседания вперемешку\n• Вступление через почту\n• Устаревший визуальный язык\n• Слабая мобильная читаемость"
    new = "• Действие и заседание сразу\n• Каталог с 6 понятными разделами\n• Календарь + архив без путаницы\n• Форма членства и понятные шаги\n• Сдержанный медицинский дизайн\n• Адаптив под телефон и проектор"
    add_textbox(s, Inches(0.85), Inches(2.4), Inches(5.2), Inches(3.8), old, size=15, color=INK)
    add_textbox(s, Inches(7.15), Inches(2.4), Inches(5.2), Inches(3.8), new, size=15, color=INK)
    add_footer(s, 4, total)

    # 5 Library screenshot
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_textbox(s, Inches(0.5), Inches(0.25), Inches(12), Inches(0.3), "БИБЛИОТЕКА", size=12, bold=True, color=COPPER)
    add_textbox(s, Inches(0.5), Inches(0.55), Inches(12), Inches(0.5), "Шесть разделов: книги, протоколы, санэпид, статьи, ФАР, видео", size=22, bold=True, color=INK, font="Georgia")
    fit_image(s, SHOTS / "02-library.png", Inches(0.8), Inches(1.2), Inches(11.7), Inches(5.5))
    add_footer(s, 5, total)

    # 6 Protocols detail
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_textbox(s, Inches(0.5), Inches(0.25), Inches(12), Inches(0.3), "ПРОТОКОЛЫ И САНЭПИД", size=12, bold=True, color=COPPER)
    add_textbox(s, Inches(0.5), Inches(0.55), Inches(12), Inches(0.5), "Документы разложены так, как ими пользуются врачи", size=22, bold=True, color=INK, font="Georgia")
    img = SHOTS / "08-library-protocols.png"
    if not img.exists():
        img = SHOTS / "02-library.png"
    fit_image(s, img, Inches(0.8), Inches(1.2), Inches(11.7), Inches(5.5))
    add_footer(s, 6, total)

    # 7 Events
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_textbox(s, Inches(0.5), Inches(0.25), Inches(12), Inches(0.3), "СОБЫТИЯ", size=12, bold=True, color=COPPER)
    add_textbox(s, Inches(0.5), Inches(0.55), Inches(12), Inches(0.5), "Календарь ближайших заседаний и архив съездов", size=22, bold=True, color=INK, font="Georgia")
    fit_image(s, SHOTS / "03-events.png", Inches(0.8), Inches(1.2), Inches(11.7), Inches(5.5))
    add_footer(s, 7, total)

    # 8 Join + patients side by side conceptually - two slides better
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_textbox(s, Inches(0.5), Inches(0.25), Inches(12), Inches(0.3), "ВСТУПЛЕНИЕ", size=12, bold=True, color=COPPER)
    add_textbox(s, Inches(0.5), Inches(0.55), Inches(12), Inches(0.5), "Путь в членство: статус → заявка → взнос", size=22, bold=True, color=INK, font="Georgia")
    fit_image(s, SHOTS / "05-join.png", Inches(0.8), Inches(1.2), Inches(11.7), Inches(5.5))
    add_footer(s, 8, total)

    # 9 Patients
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_textbox(s, Inches(0.5), Inches(0.25), Inches(12), Inches(0.3), "ПАЦИЕНТАМ", size=12, bold=True, color=COPPER)
    add_textbox(s, Inches(0.5), Inches(0.55), Inches(12), Inches(0.5), "Спокойные ответы без ложных обещаний", size=22, bold=True, color=INK, font="Georgia")
    fit_image(s, SHOTS / "04-patients.png", Inches(0.8), Inches(1.2), Inches(11.7), Inches(5.5))
    add_footer(s, 9, total)

    # 10 Cabinet
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_textbox(s, Inches(0.5), Inches(0.25), Inches(12), Inches(0.3), "КАБИНЕТ ЧЛЕНА", size=12, bold=True, color=COPPER)
    add_textbox(s, Inches(0.5), Inches(0.55), Inches(12), Inches(0.5), "Демо личного кабинета — задел под закрытую зону и чат", size=22, bold=True, color=INK, font="Georgia")
    fit_image(s, SHOTS / "06-cabinet.png", Inches(0.8), Inches(1.2), Inches(11.7), Inches(5.5))
    add_footer(s, 10, total)

    # 11 Advantages
    s = prs.slides.add_slide(blank)
    add_bg(s)
    add_textbox(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.3), "ПРЕИМУЩЕСТВА ПРОДУКТА", size=12, bold=True, color=COPPER)
    add_textbox(s, Inches(0.6), Inches(0.7), Inches(12), Inches(0.8), "Почему это лучше «просто обновить WordPress»", size=26, bold=True, color=INK, font="Georgia")
    adv = [
        ("I", "Архитектура под задачи службы АиР, а не под шаблон блога."),
        ("II", "Контент заказчика уже разложен: протоколы и санэпид из рабочих перечней."),
        ("III", "Готовность к хостингу общества: статика сейчас, сервер чата — следующим этапом."),
        ("IV", "Можно согласовывать по экранам: каждый раздел читается отдельно."),
    ]
    for i, (num, text) in enumerate(adv):
        y = Inches(1.8 + i * 1.15)
        add_textbox(s, Inches(0.7), y, Inches(1), Inches(0.5), num, size=22, bold=True, color=COPPER, font="Georgia")
        add_textbox(s, Inches(1.7), y, Inches(10.5), Inches(0.8), text, size=18, color=INK)
    add_footer(s, 11, total)

    # 12 CTA
    s = prs.slides.add_slide(blank)
    add_bg(s, PINE_DARK)
    add_textbox(s, Inches(0.7), Inches(1.5), Inches(12), Inches(0.3), "ПРЕДЛОЖЕНИЕ", size=14, bold=True, color=COPPER)
    add_textbox(s, Inches(0.7), Inches(2.1), Inches(11.5), Inches(1.8), "Согласовать макет\nи перенести на хостинг общества", size=34, bold=True, color=WHITE, font="Georgia")
    add_textbox(s, Inches(0.7), Inches(4.3), Inches(11), Inches(1.2), "Живой сайт для показа:\nhttps://leshahker.github.io/bsaer-mockup/\n\nКабинет: /cabinet/  ·  демо-вход admin / admin", size=16, color=RGBColor(0xC5, 0xD4, 0xCE))
    add_textbox(s, Inches(0.7), Inches(6.5), Inches(11), Inches(0.4), "Спасибо · БОАР · сентябрь 2026", size=12, color=MUTED)

    prs.save(OUT)
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
